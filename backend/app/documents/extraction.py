from __future__ import annotations

import glob
import io
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from app.documents.schemas import ExtractedDocument, PageText

ALLOWED_EXTENSIONS = frozenset({"pdf", "docx", "txt", "md", "png", "jpg", "jpeg", "webp"})

_IMAGE_EXTENSIONS = frozenset({"png", "jpg", "jpeg", "webp"})
_PDF_EXTENSION = "pdf"
_DOCX_EXTENSION = "docx"
_TEXT_EXTENSIONS = frozenset({"txt", "md"})

# When pdftotext returns fewer than this many characters for a page, treat the
# page as a scan/photo and fall back to OCR.
_SCANNED_PAGE_MIN_CHARS = 40


class DocumentExtractionError(Exception):
    """Raised when a document cannot be parsed/extracted."""


def require_binary(name: str) -> str:
    path = shutil.which(name)
    if path is None:
        raise DocumentExtractionError(
            f"required system binary '{name}' is not installed (poppler-utils)."
        )
    return path


def _run(argv: list[str], input_bytes: bytes | None = None) -> bytes:
    try:
        proc = subprocess.run(
            argv,
            input=input_bytes,
            capture_output=True,
            timeout=120,
            check=False,
        )
    except subprocess.TimeoutExpired as exc:  # pragma: no cover - environment bound
        raise DocumentExtractionError(f"command timed out: {' '.join(argv)}") from exc
    if proc.returncode != 0:
        raise DocumentExtractionError(
            f"command failed ({proc.returncode}): {' '.join(argv)}"
        )
    return proc.stdout


def pdf_to_pages(path: Path) -> list[PageText]:
    """Extract the text layer of a PDF, one page at a time.

    Pages with no/little text are flagged ``scanned`` so callers can OCR them.
    """
    binary = require_binary("pdftotext")
    # ``-`` writes to stdout; pages are separated by form-feed characters.
    out = _run([binary, "-layout", "-enc", "UTF-8", str(path), "-"])
    raw_pages = out.decode("utf-8", errors="replace").split("\f")
    pages: list[PageText] = []
    for index, raw in enumerate(raw_pages, start=1):
        text = raw.strip()
        pages.append(
            PageText(
                page_no=index,
                text=text,
                scanned=len(text) < _SCANNED_PAGE_MIN_CHARS,
            )
        )
    return pages or [PageText(page_no=1, text="", scanned=True)]


def render_pdf_page_jpeg(path: Path, page_no: int) -> bytes:
    """Render one PDF page as a JPEG in memory (poppler ``pdftoppm``)."""
    binary = require_binary("pdftoppm")
    with tempfile.TemporaryDirectory() as tmp:
        prefix = os.path.join(tmp, "page")
        _run(
            [
                binary,
                "-jpeg",
                "-r",
                "150",
                "-f",
                str(page_no),
                "-l",
                str(page_no),
                str(path),
                prefix,
            ]
        )
        matches = glob.glob(f"{prefix}-*.jpg")
        if not matches:
            raise DocumentExtractionError(
                f"pdftoppm produced no image for page {page_no}"
            )
        return Path(matches[0]).read_bytes()


def ocr_image(data: bytes) -> str:
    """OCR a single image (JPEG/PNG/WebP) with tesseract."""
    try:
        from PIL import Image

        import pytesseract
    except ImportError as exc:  # pragma: no cover - deps declared in pyproject
        raise DocumentExtractionError(
            "OCR support requires pytesseract and Pillow"
        ) from exc
    try:
        with Image.open(io.BytesIO(data)) as img:
            return pytesseract.image_to_string(img).strip()
    except Exception as exc:
        raise DocumentExtractionError(f"OCR failed: {exc}") from exc


def _read_image_text(path: Path, ext: str) -> list[PageText]:
    text = ocr_image(path.read_bytes())
    return [PageText(page_no=1, text=text, scanned=True)]


def _read_docx_text(path: Path) -> list[PageText]:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - deps declared in pyproject
        raise DocumentExtractionError("docx support requires python-docx") from exc
    document = docx.Document(str(path))
    chunks: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            chunks.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return [PageText(page_no=1, text="\n".join(chunks), scanned=False)]


def _read_text_file(path: Path) -> list[PageText]:
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    return [PageText(page_no=1, text=text, scanned=False)]


def _read_pdf_text(path: Path) -> list[PageText]:
    pages = pdf_to_pages(path)
    ocr_missing = [page for page in pages if page.scanned]
    if not ocr_missing:
        return pages
    for page in ocr_missing:
        try:
            jpeg = render_pdf_page_jpeg(path, page.page_no)
            page.text = ocr_image(jpeg) or page.text
            page.scanned = True
        except DocumentExtractionError:
            continue  # keep empty page; OCR is best-effort
    return pages

def extract_text(path: Path, file_name: str) -> ExtractedDocument:
    """Run the textract-style extraction pipeline for a stored upload."""
    ext = file_name.rsplit(".", maxsplit=1)[-1].lower() if "." in file_name else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise DocumentExtractionError(f"unsupported file extension '.{ext}'")
    if ext == _PDF_EXTENSION:
        pages = _read_pdf_text(path)
    elif ext == _DOCX_EXTENSION:
        pages = _read_docx_text(path)
    elif ext in _TEXT_EXTENSIONS:
        pages = _read_text_file(path)
    elif ext in _IMAGE_EXTENSIONS:
        pages = _read_image_text(path, ext)
    else:  # pragma: no cover - guarded above
        raise DocumentExtractionError(f"unsupported file extension '.{ext}'")
    return ExtractedDocument(file_name=file_name, kind=ext, pages=pages)


def render_visual_pages(path: Path, ext: str, max_pages: int) -> list[tuple[int, bytes]]:
    """Return (page_no, jpeg_bytes) candidates for visual analysis.

    PDFs render up to ``max_pages`` pages; a standalone image is returned as-is
    (converted to JPEG). Other kinds yield nothing.
    """
    if ext in _IMAGE_EXTENSIONS:
        data = path.read_bytes()
        if ext != "jpeg" and ext != "jpg":
            try:
                from PIL import Image

                with Image.open(io.BytesIO(data)) as img:
                    if img.mode in ("RGBA", "P", "LA"):
                        img = img.convert("RGB")
                    buf = io.BytesIO()
                    img.save(buf, format="JPEG", quality=90)
                    data = buf.getvalue()
            except Exception:
                pass  # send original; the vision endpoint validates formats
        return [(1, data)]
    if ext == _PDF_EXTENSION:
        out: list[tuple[int, bytes]] = []
        for page_no in range(1, max_pages + 1):
            try:
                out.append((page_no, render_pdf_page_jpeg(path, page_no)))
            except DocumentExtractionError:
                break
        return out
    return []
